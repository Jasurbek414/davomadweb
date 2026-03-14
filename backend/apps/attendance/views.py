from rest_framework import viewsets, generics
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter, OrderingFilter
from django.db.models import Count, Q
from django.http import HttpResponse
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from .models import AttendanceRecord
from .serializers import AttendanceRecordSerializer
from apps.accounts.permissions import IsTeacherOrAbove, IsOperatorOrAbove


class AttendanceRecordViewSet(viewsets.ModelViewSet):
    queryset = AttendanceRecord.objects.select_related(
        'student', 'student__class_ref', 'student__school',
        'check_in_device', 'check_out_device'
    ).order_by('-date', 'student__last_name')
    serializer_class = AttendanceRecordSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ['date', 'status', 'student__school', 'student__class_ref', 'student']
    search_fields = ['student__first_name', 'student__last_name', 'student__student_id']
    ordering_fields = ['date', 'check_in', 'status']

    def get_permissions(self):
        if self.action in ['create', 'update', 'partial_update']:
            return [IsOperatorOrAbove()]
        return [IsTeacherOrAbove()]

    def get_queryset(self):
        user = self.request.user
        qs = super().get_queryset()
        if user.is_superuser or user.has_role('superadmin'):
            return qs
        if user.has_role('region_director'):
            region_ids = user.user_roles.filter(role__name='region_director', is_active=True).values_list('region_id', flat=True)
            return qs.filter(student__school__district__region_id__in=region_ids)
        if user.has_role('district_director'):
            district_ids = user.user_roles.filter(role__name='district_director', is_active=True).values_list('district_id', flat=True)
            return qs.filter(student__school__district_id__in=district_ids)
        school_ids = user.user_roles.filter(
            role__name__in=['school_director', 'operator', 'teacher'],
            is_active=True, school__isnull=False
        ).values_list('school_id', flat=True)
        if school_ids:
            return qs.filter(student__school_id__in=school_ids)
        # Parent: own children only
        if user.has_role('parent'):
            try:
                parent = user.parent_profile
                student_ids = parent.student_links.filter(status='active').values_list('student_id', flat=True)
                return qs.filter(student_id__in=student_ids)
            except Exception:
                pass
        return qs.none()

    @action(detail=False, methods=['get'], url_path='statistics')
    def statistics(self, request):
        """Davomad statistikasi"""
        school_id = request.query_params.get('school')
        class_id = request.query_params.get('class_ref')
        start_date = request.query_params.get('start_date', str(date.today()))
        end_date = request.query_params.get('end_date', str(date.today()))

        qs = self.get_queryset().filter(date__range=[start_date, end_date])
        if school_id:
            qs = qs.filter(student__school_id=school_id)
        if class_id:
            qs = qs.filter(student__class_ref_id=class_id)

        stats = qs.aggregate(
            total=Count('id'),
            present=Count('id', filter=Q(status='present')),
            late=Count('id', filter=Q(status='late')),
            absent=Count('id', filter=Q(status='absent')),
            excused=Count('id', filter=Q(status='excused')),
        )
        total = stats['total'] or 1
        stats['attendance_rate'] = round(
            ((stats['present'] + stats['late']) / total) * 100, 1
        )
        return Response(stats)

    @action(detail=False, methods=['get'], url_path='today')
    def today(self, request):
        today = date.today()
        qs = self.get_queryset().filter(date=today)
        school_id = request.query_params.get('school')
        if school_id:
            qs = qs.filter(student__school_id=school_id)
        serializer = self.get_serializer(qs, many=True, context={'request': request})
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """Davomadni Excel faylga yuklash"""
        qs = self.filter_queryset(self.get_queryset())
        
        # Create workbook and sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Davomad"
        
        # Header row
        headers = [
            '№', 'Sana', 'O\'quvchi F.I.O', 'Sinf', 
            'ID raqam', 'Kelish vaqti', 'Ketish vaqti', 
            'Holat', 'Kechikish (daqiqa)'
        ]
        
        # Styling
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            
        # Data rows
        status_map = {
            'present': 'Keldi',
            'late': 'Kechikdi',
            'absent': 'Kelmadi',
            'excused': 'Sababli',
        }
        
        for idx, record in enumerate(qs, 1):
            row_num = idx + 1
            ws.cell(row=row_num, column=1, value=idx)
            ws.cell(row=row_num, column=2, value=str(record.date))
            ws.cell(row=row_num, column=3, value=record.student.full_name)
            ws.cell(row=row_num, column=4, value=record.student.class_ref.name if record.student.class_ref else '-')
            ws.cell(row=row_num, column=5, value=record.student.student_id)
            ws.cell(row=row_num, column=6, value=record.check_in.strftime('%H:%M:%S') if record.check_in else '-')
            ws.cell(row=row_num, column=7, value=record.check_out.strftime('%H:%M:%S') if record.check_out else '-')
            ws.cell(row=row_num, column=8, value=status_map.get(record.status, record.status))
            ws.cell(row=row_num, column=9, value=record.late_minutes if record.late_minutes else 0)
            
        # Column widths
        dims = {}
        for row in ws.rows:
            for cell in row:
                if cell.value:
                    dims[cell.column_letter] = max((dims.get(cell.column_letter, 0), len(str(cell.value))))
        for col, value in dims.items():
            ws.column_dimensions[col].width = value + 2

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename=davomad_{date.today()}.xlsx'
        wb.save(response)
        return response
